import docx

doc = docx.Document(r'C:\Users\Andndre\Downloads\Blueprint-SmartPrasada.docx')
print(f'Total paragraphs: {len(doc.paragraphs)}')

targets = [
    '2.2 Alur inti',
    '2.3 Alur e-learning',
    '2.4 Alur fasilitator',
    '2.5 Alur pengelolaan konten',
    '2.6 Alur data penelitian',
    '3.1 Arsitektur lima modul',
    '3.3 Arsitektur aplikasi',
    '3.4 Susunan runtime VR',
    '5.2 Peta navigasi'
]

for idx, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    for t in targets:
        if t.lower() in txt.lower():
            print(f'[{idx}] ({t}) -> {txt[:80]}')
