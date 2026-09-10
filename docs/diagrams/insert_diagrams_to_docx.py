import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image
import os, sys

if sys.stdout.encoding != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8')

source_path = r'C:\Users\Andndre\Downloads\Blueprint-SmartPrasada.docx'
target_downloads = r'C:\Users\Andndre\Downloads\Blueprint-SmartPrasada-With-Diagrams.docx'
target_project = r'D:\smart_prasada\docs\Blueprint-SmartPrasada-With-Diagrams.docx'
diagrams_dir = r'D:\smart_prasada\docs\diagrams'

doc = docx.Document(source_path)

# Map diagrams to their target section titles and figure captions
diagram_configs = [
    {
        'key': '2.2 Alur inti: sesi Virtual Living Museum',
        'match_text': '2.2 Alur inti: sesi Virtual Living Museum',
        'sub_match': 'ORIENTASI',
        'img': os.path.join(diagrams_dir, '01_alur_inti_sesi_vr_4_fase.png'),
        'caption': 'Gambar 2.1 Alur Inti Sesi Virtual Living Museum (Empat Fase)',
        'max_height': 8.6,
        'max_width': 5.5
    },
    {
        'key': '2.3 Alur e-learning berjenjang',
        'match_text': '2.3 Alur e-learning berjenjang',
        'sub_match': 'PRE-TEST (1)',
        'img': os.path.join(diagrams_dir, '02_alur_elearning_berjenjang.png'),
        'caption': 'Gambar 2.2 Alur E-Learning Berjenjang dan Gating Pembelajaran',
        'max_height': 8.5,
        'max_width': 5.8
    },
    {
        'key': '2.4 Alur fasilitator (mode kiosk',
        'match_text': '2.4 Alur fasilitator (mode kiosk',
        'sub_match': 'Catatan operasional: parameter kiosk',
        'img': os.path.join(diagrams_dir, '03_alur_fasilitator_kiosk.png'),
        'caption': 'Gambar 2.3 Alur Operasional Fasilitator, Peluncur Sesi QR, dan Mode Kiosk',
        'max_height': 8.5,
        'max_width': 5.5
    },
    {
        'key': '2.5 Alur pengelolaan konten (administrator)',
        'match_text': '2.5 Alur pengelolaan konten (administrator)',
        'sub_match': '2.5 Alur pengelolaan konten (administrator)',
        'img': os.path.join(diagrams_dir, '04_alur_pengelolaan_konten_admin.png'),
        'caption': 'Gambar 2.4 Alur Tujuh Tahap Pengelolaan Konten Pembelajaran dan Museum Virtual',
        'max_height': 8.5,
        'max_width': 5.5
    },
    {
        'key': '2.6 Alur data penelitian',
        'match_text': '2.6 Alur data penelitian',
        'sub_match': 'Kuesioner keterpakaian pasca-sesi',
        'img': os.path.join(diagrams_dir, '05_alur_data_penelitian_tkt.png'),
        'caption': 'Gambar 2.5 Pipeline Data Penelitian Lintas Instrumen dan Validasi TKT 6',
        'max_height': 8.5,
        'max_width': 5.5
    },
    {
        'key': '3.1 Arsitektur lima modul',
        'match_text': '3.1 Arsitektur lima modul',
        'sub_match': 'Kolom terakhir menunjukkan realisasinya dalam kode',
        'img': os.path.join(diagrams_dir, '06_arsitektur_lima_modul_sistem.png'),
        'caption': 'Gambar 3.1 Arsitektur Sistem Lima Modul SmartPrasada dan Lapisan Pendukung',
        'max_height': 8.5,
        'max_width': 5.5
    },
    {
        'key': '3.4 Susunan runtime VR',
        'match_text': '3.4 Susunan runtime VR',
        'sub_match': 'dapat diverifikasi tanpa headset',
        'img': os.path.join(diagrams_dir, '07_arsitektur_runtime_vr_dan_auth_token.png'),
        'caption': 'Gambar 3.2 Arsitektur Modular Runtime VR dan Autentikasi Stateless Token HMAC',
        'max_height': 8.5,
        'max_width': 5.5
    },
    {
        'key': '5.2 Peta navigasi',
        'match_text': '5.2 Peta navigasi',
        'sub_match': '└── Umpan Balik',
        'img': os.path.join(diagrams_dir, '08_peta_navigasi_aplikasi.png'),
        'caption': 'Gambar 5.1 Peta Navigasi dan Hierarki Antarmuka Sistem SmartPrasada',
        'max_height': 8.5,
        'max_width': 5.5
    }
]

def calculate_dimensions(img_path, max_w_inch, max_h_inch):
    with Image.open(img_path) as im:
        w_px, h_px = im.size
    aspect = w_px / h_px
    target_w = max_w_inch
    target_h = target_w / aspect
    if target_h > max_h_inch:
        target_h = max_h_inch
        target_w = target_h * aspect
    return Inches(target_w), Inches(target_h)

inserted_count = 0

# Find paragraph targets (skipping Table of Contents which is in the first 50 paragraphs)
for cfg in diagram_configs:
    target_idx = None
    for idx, p in enumerate(doc.paragraphs):
        if idx < 50:
            continue
        if cfg['sub_match'] in p.text:
            target_idx = idx
            break
        elif cfg['match_text'] in p.text and target_idx is None:
            target_idx = idx

    if target_idx is not None:
        target_p = doc.paragraphs[target_idx]
        print(f"Matched '{cfg['caption']}' at paragraph [{target_idx}]: {target_p.text[:60]}")

        # Compute optimal display size
        w_inch, h_inch = calculate_dimensions(cfg['img'], cfg['max_width'], cfg['max_height'])

        # Create paragraph for image
        img_p = doc.add_paragraph()
        target_p._element.addnext(img_p._element)
        img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = img_p.add_run()
        run.add_picture(cfg['img'], width=w_inch, height=h_inch)

        # Create paragraph for caption
        cap_p = doc.add_paragraph()
        img_p._element.addnext(cap_p._element)
        cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_run = cap_p.add_run(cfg['caption'])
        cap_run.bold = True
        cap_run.font.size = Pt(10)
        cap_run.font.color.rgb = RGBColor(0x33, 0x41, 0x55) # Slate 700

        # Small spacing after caption
        cap_p.paragraph_format.space_after = Pt(12)
        cap_p.paragraph_format.space_before = Pt(4)

        inserted_count += 1
    else:
        print(f"FAILED to match target for: {cfg['caption']}")

print(f"\nTotal diagrams inserted: {inserted_count}/8")

# Save duplicates
doc.save(target_downloads)
print(f"Saved duplicate to: {target_downloads}")

doc.save(target_project)
print(f"Saved duplicate to: {target_project}")
